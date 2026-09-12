"""
Purpose: Excel Export Service
Author: Scheduler Backend Team
Created Date: 2026-08-03
Module Description: Generates a styled multi-sheet Excel workbook (one sheet
per class, a consolidated faculty sheet, and a statistics sheet) from the
saved timetable records.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.services.export_data import load_export_data
from app.services.statistics import StatisticsService
from app.utils.logger import get_logger

logger = get_logger(__name__)

FILLS = {
    "Theory": PatternFill(start_color="EBF3FC", end_color="EBF3FC", fill_type="solid"),
    "Lab": PatternFill(start_color="EAF9EB", end_color="EAF9EB", fill_type="solid"),
    "Elective": PatternFill(start_color="FEF5E7", end_color="FEF5E7", fill_type="solid"),
    "Free": PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"),
    "Break": PatternFill(start_color="F2F4F5", end_color="F2F4F5", fill_type="solid"),
    "Lunch": PatternFill(start_color="E6E8EA", end_color="E6E8EA", fill_type="solid"),
}


class ExcelExportService:
    """
    Exports the generated timetable to a styled Excel workbook.
    """

    @classmethod
    async def export_timetable(
        cls,
        db,
        output_path: Path,
        user_id: Optional[str] = None,
    ) -> None:
        """
        Query timetable entries, organize them by Class and Faculty, and write
        a multi-sheet spreadsheet.

        Args:
            db (AsyncSession): Active database session.
            output_path (Path): Destination path for the saved Excel.
            user_id (Optional[str]): Scopes the Setup config lookup.
        """
        logger.info(f"Generating Excel timetable export at: {output_path}")

        class_models, config, formatted = await load_export_data(db, user_id)
        working_days = config["working_days"]
        periods = config["periods"]

        wb = Workbook()
        default_sheet = wb.active
        wb.remove(default_sheet)

        font_family = "Segoe UI"
        header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
        title_font = Font(name=font_family, size=15, bold=True, color="1B365D")
        body_font = Font(name=font_family, size=10)
        bold_body_font = Font(name=font_family, size=10, bold=True)
        italic_body_font = Font(name=font_family, size=9, italic=True, color="555555")

        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")

        thin_side = Side(border_style="thin", color="CCCCCC")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Build period headers from the interleaved periods array
        period_headers = []
        for p in periods:
            if p.get("is_break", False):
                period_headers.append("Break")
            elif p.get("is_lunch", False):
                period_headers.append("Lunch")
            else:
                period_headers.append(f"Period {p['period_number']}")

        # 1. One sheet per class
        used_titles = set()
        for cm in class_models:
            label = cm["label"]
            title = f"{cm['short_code']} - Section {cm['section']}"[:31]
            unique_title = title
            suffix = 2
            while unique_title in used_titles:
                unique_title = f"{title[:27]} ({suffix})"
                suffix += 1
            used_titles.add(unique_title)
            ws = wb.create_sheet(title=unique_title)
            ws.views.sheetView[0].showGridLines = True

            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(period_headers))
            ws.cell(row=1, column=1, value=f"{cm['short_code']} - Section {cm['section']}  |  Room {cm.get('room_no') or '-'}").font = title_font
            ws.cell(row=1, column=1).alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 36

            ws.row_dimensions[3].height = 30
            for col_idx, h in enumerate(["Day"] + period_headers, 1):
                cell = ws.cell(row=3, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border

            r_idx = 4
            for day in working_days:
                ws.row_dimensions[r_idx].height = 40
                day_cell = ws.cell(row=r_idx, column=1, value=day)
                day_cell.font = bold_body_font
                day_cell.alignment = center_align
                day_cell.border = thin_border

                day_slots = formatted.get(label, {}).get(day, [])
                for idx, slot in enumerate(day_slots):
                    col_idx = idx + 2
                    cell = ws.cell(row=r_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = center_align
                    slot_type = slot["type"]
                    cell.fill = FILLS.get(slot_type, FILLS["Free"])

                    if slot_type in ("Break", "Lunch"):
                        cell.value = slot["subject"].upper()
                        cell.font = italic_body_font
                    elif slot_type == "Free":
                        cell.value = "Free Period"
                        cell.font = italic_body_font
                    else:
                        cell.value = f"{slot['subject']}\n({slot['faculty']})"
                        cell.font = body_font
                r_idx += 1

            ws.column_dimensions["A"].width = 14
            for col_idx in range(2, len(periods) + 2):
                ws.column_dimensions[get_column_letter(col_idx)].width = 22

        # 2. Consolidated faculty sheet
        ws_fac = wb.create_sheet(title="Faculty Timetable")
        ws_fac.views.sheetView[0].showGridLines = True

        ws_fac.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2 + len(period_headers))
        ws_fac.cell(row=1, column=1, value="Consolidated Faculty Timetable").font = title_font
        ws_fac.row_dimensions[1].height = 36

        ws_fac.row_dimensions[3].height = 28
        for col_idx, h in enumerate(["Faculty", "Day"] + period_headers, 1):
            cell = ws_fac.cell(row=3, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        faculty_schedule = cls._build_faculty_schedule(formatted)
        faculty_names = sorted(faculty_schedule.keys())

        r_idx = 4
        for fac_name in faculty_names:
            start_row = r_idx
            for day_idx, day in enumerate(working_days):
                ws_fac.row_dimensions[r_idx].height = 32

                if day_idx == 0:
                    fac_cell = ws_fac.cell(row=r_idx, column=1, value=fac_name)
                    fac_cell.font = bold_body_font
                else:
                    fac_cell = ws_fac.cell(row=r_idx, column=1)
                fac_cell.alignment = center_align
                fac_cell.border = thin_border

                day_cell = ws_fac.cell(row=r_idx, column=2, value=day)
                day_cell.font = body_font
                day_cell.alignment = center_align
                day_cell.border = thin_border

                for idx, p in enumerate(periods):
                    col_idx = idx + 3
                    cell = ws_fac.cell(row=r_idx, column=col_idx)
                    cell.border = thin_border
                    cell.alignment = center_align
                    if p.get("is_lunch", False):
                        cell.value = "LUNCH BREAK"
                        cell.font = italic_body_font
                        cell.fill = FILLS["Lunch"]
                    elif p.get("is_break", False):
                        cell.value = "BREAK"
                        cell.font = italic_body_font
                        cell.fill = FILLS["Break"]
                    else:
                        p_num = p["period_number"]
                        slot = faculty_schedule.get(fac_name, {}).get(day, {}).get(p_num)
                        if slot:
                            cell.value = f"{slot['class']}\n({slot['subject']})"
                            cell.font = body_font
                            cell.fill = FILLS.get(slot["type"], FILLS["Theory"])
                        else:
                            cell.value = "-"
                            cell.font = italic_body_font
                            cell.fill = FILLS["Free"]
                r_idx += 1
            ws_fac.merge_cells(start_row=start_row, start_column=1, end_row=r_idx - 1, end_column=1)

        ws_fac.column_dimensions["A"].width = 22
        ws_fac.column_dimensions["B"].width = 12
        for col_idx in range(3, len(periods) + 3):
            ws_fac.column_dimensions[get_column_letter(col_idx)].width = 20

        # 3. Statistics sheet
        ws_stats = wb.create_sheet(title="Statistics")
        ws_stats.views.sheetView[0].showGridLines = True

        ws_stats.merge_cells("A1:E1")
        ws_stats["A1"] = "Timetable System Analytics"
        ws_stats["A1"].font = title_font
        ws_stats.row_dimensions[1].height = 36

        stats_payload = await StatisticsService.get_faculty_statistics(db, user_id=user_id)
        summary = stats_payload["summary"]
        fac_stats_list = stats_payload["faculty_statistics"]

        ws_stats.cell(row=3, column=1, value="Analytical Metric").font = header_font
        ws_stats.cell(row=3, column=1).fill = header_fill
        ws_stats.cell(row=3, column=1).border = thin_border
        ws_stats.cell(row=3, column=2, value="Value").font = header_font
        ws_stats.cell(row=3, column=2).fill = header_fill
        ws_stats.cell(row=3, column=2).border = thin_border
        ws_stats.cell(row=3, column=2).alignment = center_align

        metrics = [
            ("Total Registered Classes", summary["total_classes"]),
            ("Total Registered Subjects", summary["total_subjects"]),
            ("Total Registered Faculty", summary["total_faculty"]),
            ("Total Scheduled Teaching Periods", summary["total_scheduled_periods"]),
            ("Overall Timetable Utilization (%)", f"{summary['overall_utilization_percentage']}%"),
        ]

        r = 4
        for metric, val in metrics:
            m_cell = ws_stats.cell(row=r, column=1, value=metric)
            m_cell.font = bold_body_font
            m_cell.border = thin_border
            v_cell = ws_stats.cell(row=r, column=2, value=val)
            v_cell.font = body_font
            v_cell.border = thin_border
            v_cell.alignment = center_align
            r += 1

        r += 2
        ws_stats.cell(row=r, column=1, value="Faculty Teaching Hours Summary").font = title_font
        r += 1

        fac_headers = ["Faculty Name", "Subject", "Weekly Scheduled Periods", "Classes Handled"]
        ws_stats.row_dimensions[r].height = 25
        for idx, h in enumerate(fac_headers, 1):
            cell = ws_stats.cell(row=r, column=idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = center_align

        r += 1
        for fs in fac_stats_list:
            ws_stats.row_dimensions[r].height = 20
            ws_stats.cell(row=r, column=1, value=fs["faculty_name"]).font = body_font
            ws_stats.cell(row=r, column=1).border = thin_border
            ws_stats.cell(row=r, column=2, value=fs.get("subject") or "-").font = body_font
            ws_stats.cell(row=r, column=2).border = thin_border
            hours_cell = ws_stats.cell(row=r, column=3, value=fs["total_hours"])
            hours_cell.font = bold_body_font
            hours_cell.border = thin_border
            hours_cell.alignment = center_align
            ws_stats.cell(row=r, column=4, value=", ".join(fs["classes_taught"])).font = body_font
            ws_stats.cell(row=r, column=4).border = thin_border
            r += 1

        ws_stats.column_dimensions["A"].width = 30
        ws_stats.column_dimensions["B"].width = 28
        ws_stats.column_dimensions["C"].width = 26
        ws_stats.column_dimensions["D"].width = 40

        wb.save(output_path)
        logger.info("Successfully exported Excel workbook.")

    @classmethod
    def _build_faculty_schedule(cls, formatted: Dict[str, Any]) -> Dict[str, Any]:
        """
        Rebuild the faculty -> day -> period mapping from the formatted class
        timetable.
        """
        schedule: Dict[str, Any] = {}
        for label, days in formatted.items():
            for day, slots in days.items():
                for slot in slots:
                    if slot["type"] in ("Break", "Lunch", "Free"):
                        continue
                    schedule.setdefault(slot["faculty"], {}).setdefault(day, {})[slot["period"]] = {
                        "class": label,
                        "subject": slot["subject"],
                        "type": slot["type"],
                    }
        return schedule
