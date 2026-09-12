"""
Purpose: Word (.docx) Export Service
Author: Smart Timetable Backend Team
Module Description: Generates a Word document containing each class's weekly
timetable as a table (Day x Period), with a header showing the class name,
short code, section and room. Uses python-docx.
"""

from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.services.export_data import load_export_data
from app.utils.logger import get_logger

logger = get_logger(__name__)


class WordExportService:
    """
    Exports the generated timetable to a Word (.docx) document.
    """

    @classmethod
    async def export_timetable(
        cls,
        db,
        output_path: Path,
        user_id: Optional[str] = None,
    ) -> None:
        """
        Create a .docx file from the saved timetable records.

        Args:
            db (AsyncSession): Active database session.
            output_path (Path): File path destination.
            user_id (Optional[str]): Scopes the Setup config lookup.
        """
        logger.info(f"Generating Word timetable export at: {output_path}")

        class_models, config, formatted = await load_export_data(db, user_id)
        working_days = config["working_days"]
        periods = config["periods"]

        if not formatted:
            raise ValueError("No timetable data available for export.")

        doc = Document()

        section = doc.sections[0]
        section.page_width = Pt(842)
        section.page_height = Pt(595)
        section.left_margin = Pt(40)
        section.right_margin = Pt(40)
        section.top_margin = Pt(40)
        section.bottom_margin = Pt(40)

        styles = doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10)

        for cm in class_models:
            label = cm["label"]
            class_data = formatted.get(label, {})

            meta = doc.add_table(rows=1, cols=4)
            meta.style = "Table Grid"
            meta.alignment = WD_TABLE_ALIGNMENT.CENTER
            meta_cells = meta.rows[0].cells
            meta_cells[0].text = f"Branch: {cm['short_code']}"
            meta_cells[1].text = f"Section: {cm['section']}"
            meta_cells[2].text = f"Room: {cm.get('room_no') or '-'}"
            meta_cells[3].text = "Academic Period 2026"

            title = doc.add_paragraph()
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_run = title.add_run(cm["class_name"])
            title_run.bold = True
            title_run.font.size = Pt(18)
            title_run.font.color.rgb = cls._rgb("1B365D")

            subtitle = doc.add_paragraph()
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle_run = subtitle.add_run("TIME TABLE")
            subtitle_run.bold = True
            subtitle_run.font.size = Pt(14)
            subtitle_run.font.color.rgb = cls._rgb("1B365D")

            table = doc.add_table(rows=1 + len(working_days), cols=1 + len(periods))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            header_cells = table.rows[0].cells
            header_cells[0].text = "Day"
            for idx, p in enumerate(periods, start=1):
                if p.get("is_break", False):
                    header_text = f"BREAK\n{p['start_time']}-{p['end_time']}"
                elif p.get("is_lunch", False):
                    header_text = f"LUNCH\n{p['start_time']}-{p['end_time']}"
                else:
                    header_text = f"P{p['period_number']}\n{p['start_time']}-{p['end_time']}"
                header_cells[idx].text = header_text

            for row_idx, day in enumerate(working_days, start=1):
                cells = table.rows[row_idx].cells
                cells[0].text = day
                day_slots = class_data.get(day, [])
                for idx, slot in enumerate(day_slots):
                    col_idx = idx + 1
                    if slot["type"] in ("Break", "Lunch"):
                        cells[col_idx].text = slot["subject"].upper()
                    elif slot["type"] == "Free":
                        cells[col_idx].text = "Free Period"
                    else:
                        cells[col_idx].text = f"{slot['subject']}\n{slot['faculty']}"

            doc.add_page_break()

        doc.save(output_path)
        logger.info("Successfully generated Word timetable document.")

    @staticmethod
    def _rgb(hex_color: str):
        from docx.shared import RGBColor
        hex_color = hex_color.lstrip("#")
        return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))
